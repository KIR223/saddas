import io
import os
import tempfile
import unittest

from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from app.main import app


def build_upload() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "261-266"
    sheet["A1"] = "Расписание занятий на 1 семестр 2026-2027"
    sheet["C2"] = "ТМ-261"
    sheet["A3"] = "Понедельник"
    sheet.merge_cells("A3:A4")
    sheet["B3"] = 1
    sheet.merge_cells("B3:B4")
    sheet["C4"] = "Физика\nКурницкая Т.С."
    sheet["C4"].fill = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")
    sheet["D4"] = 307
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


class ApiTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        os.environ["DATABASE_PATH"] = os.path.join(self.tmp.name, "schedule.db")
        os.environ["ADMIN_TOKEN"] = "secret"
        self.client = TestClient(app)

    def tearDown(self):
        self.tmp.cleanup()

    def test_upload_and_schedule(self):
        info = self.client.get("/api/info")
        self.assertEqual(info.status_code, 200)
        payload = info.json()
        self.assertIn(payload["week"], ("green", "red"))
        self.assertEqual(payload["title"], "")
        self.assertIsNone(payload["uploaded_at"])

        self.assertEqual(self.client.get("/api/groups").json(), {"groups": []})

        denied = self.client.post(
            "/api/admin/upload",
            files={"file": ("schedule.xlsx", build_upload(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(denied.status_code, 401)

        uploaded = self.client.post(
            "/api/admin/upload",
            headers={"X-Admin-Token": "secret"},
            files={"file": ("schedule.xlsx", build_upload(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        body = uploaded.json()
        self.assertEqual(body["groups"], 1)
        self.assertEqual(body["lessons"], 1)
        self.assertEqual(body["warnings"], [])

        groups = self.client.get("/api/groups").json()["groups"]
        self.assertEqual(groups, ["ТМ-261"])

        schedule = self.client.get("/api/schedule", params={"group": "ТМ-261", "week": "green"})
        self.assertEqual(schedule.status_code, 200)
        day = schedule.json()["days"][0]
        self.assertEqual(day["day"], "Понедельник")
        self.assertEqual(day["pairs"][0]["subject"], "Физика")
        self.assertEqual(day["pairs"][0]["teacher"], "Курницкая Т.С.")
        self.assertEqual(day["pairs"][0]["room"], "307")

        red = self.client.get("/api/schedule", params={"group": "ТМ-261", "week": "red"})
        self.assertEqual(red.status_code, 200)
        self.assertEqual(red.json()["days"], [])

        missing = self.client.get("/api/schedule", params={"group": "НЕТ-000", "week": "green"})
        self.assertEqual(missing.status_code, 404)

        bad_week = self.client.get("/api/schedule", params={"group": "ТМ-261", "week": "blue"})
        self.assertEqual(bad_week.status_code, 400)

        info_after = self.client.get("/api/info").json()
        self.assertIn("Расписание занятий", info_after["title"])
        self.assertIsNotNone(info_after["uploaded_at"])

        self.assertEqual(self.client.get("/api/rooms").json(), {"rooms": ["307"]})
        self.assertEqual(self.client.get("/api/teachers").json(), {"teachers": ["Курницкая Т.С."]})

    def auth(self):
        return {"X-Admin-Token": "secret"}

    def upload_default(self):
        response = self.client.post(
            "/api/admin/upload",
            headers=self.auth(),
            files={"file": ("schedule.xlsx", build_upload(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def test_bells_regular_and_short_day(self):
        denied = self.client.put("/api/admin/bells", json={"pairs": []})
        self.assertEqual(denied.status_code, 401)

        regular = [
            {"pair": 1, "start": "08:30", "end": "10:00"},
            {"pair": 2, "start": "10:10", "end": "11:40"},
        ]
        saved = self.client.put("/api/admin/bells", headers=self.auth(), json={"pairs": regular})
        self.assertEqual(saved.status_code, 200, saved.text)
        self.assertEqual(saved.json()["kind"], "regular")
        self.assertEqual(saved.json()["pairs"], regular)

        overlap = self.client.put(
            "/api/admin/bells",
            headers=self.auth(),
            json={"pairs": [
                {"pair": 1, "start": "08:30", "end": "10:00"},
                {"pair": 2, "start": "09:50", "end": "11:00"},
            ]},
        )
        self.assertEqual(overlap.status_code, 400)

        short_day = "2026-10-08"
        short = [{"pair": 1, "start": "08:30", "end": "09:05"}]
        short_saved = self.client.put(
            f"/api/admin/bells/{short_day}",
            headers=self.auth(),
            json={"pairs": short},
        )
        self.assertEqual(short_saved.status_code, 200, short_saved.text)
        self.assertEqual(short_saved.json()["kind"], "short")
        self.assertEqual(short_saved.json()["pairs"], short)

        by_date = self.client.get("/api/bells", params={"date": short_day})
        self.assertEqual(by_date.json()["kind"], "short")
        other = self.client.get("/api/bells", params={"date": "2026-10-09"})
        self.assertEqual(other.json()["kind"], "regular")
        self.assertEqual(other.json()["pairs"], regular)

        removed = self.client.delete(f"/api/admin/bells/{short_day}", headers=self.auth())
        self.assertEqual(removed.status_code, 204)
        after = self.client.get("/api/bells", params={"date": short_day})
        self.assertEqual(after.json()["kind"], "regular")
        self.assertEqual(after.json()["pairs"], regular)

    def test_pair_change_stays_on_one_date(self):
        self.upload_default()
        changed = "2026-10-05"
        same_weekday = "2026-10-19"

        denied = self.client.put(
            f"/api/admin/schedule/{changed}/pairs/1",
            params={"group": "ТМ-261"},
            json={"subject": "Химия", "teacher": "Петров П.П.", "room": "101"},
        )
        self.assertEqual(denied.status_code, 401)

        updated = self.client.put(
            f"/api/admin/schedule/{changed}/pairs/1",
            headers=self.auth(),
            params={"group": "ТМ-261"},
            json={"subject": "Химия", "teacher": "Петров П.П.", "room": "101"},
        )
        self.assertEqual(updated.status_code, 200, updated.text)
        body = updated.json()
        self.assertTrue(body["custom"])
        self.assertEqual(body["day"], "Понедельник")
        self.assertEqual(body["week"], "green")
        self.assertEqual(body["pairs"][0]["subject"], "Химия")

        later = self.client.get("/api/schedule/day", params={"group": "ТМ-261", "date": same_weekday})
        self.assertEqual(later.status_code, 200, later.text)
        self.assertFalse(later.json()["custom"])
        self.assertEqual(later.json()["pairs"][0]["subject"], "Физика")

        template = self.client.get("/api/schedule", params={"group": "ТМ-261", "week": "green"})
        self.assertEqual(template.json()["days"][0]["pairs"][0]["subject"], "Физика")

        self.upload_default()
        still = self.client.get("/api/schedule/day", params={"group": "ТМ-261", "date": changed})
        self.assertEqual(still.json()["pairs"][0]["subject"], "Химия")

        cleared = self.client.delete(
            f"/api/admin/schedule/{changed}/pairs/1",
            headers=self.auth(),
            params={"group": "ТМ-261"},
        )
        self.assertEqual(cleared.status_code, 204)
        restored = self.client.get("/api/schedule/day", params={"group": "ТМ-261", "date": changed})
        self.assertFalse(restored.json()["custom"])
        self.assertEqual(restored.json()["pairs"][0]["subject"], "Физика")

    def test_replace_whole_day(self):
        self.upload_default()
        on = "2026-10-05"
        replaced = self.client.put(
            f"/api/admin/schedule/{on}",
            headers=self.auth(),
            params={"group": "ТМ-261"},
            json={"pairs": [{"pair": 2, "subject": "История", "teacher": "Бурич П.А.", "room": "119"}]},
        )
        self.assertEqual(replaced.status_code, 200, replaced.text)
        pairs = replaced.json()["pairs"]
        self.assertEqual(len(pairs), 1)
        self.assertEqual(pairs[0]["pair"], 2)
        self.assertEqual(pairs[0]["subject"], "История")

        other = self.client.get("/api/schedule/day", params={"group": "ТМ-261", "date": "2026-10-19"})
        self.assertEqual(other.json()["pairs"][0]["subject"], "Физика")

        removed = self.client.delete(
            f"/api/admin/schedule/{on}",
            headers=self.auth(),
            params={"group": "ТМ-261"},
        )
        self.assertEqual(removed.status_code, 204)
        back = self.client.get("/api/schedule/day", params={"group": "ТМ-261", "date": on})
        self.assertEqual(back.json()["pairs"][0]["subject"], "Физика")
        self.assertFalse(back.json()["custom"])


if __name__ == "__main__":
    unittest.main()
