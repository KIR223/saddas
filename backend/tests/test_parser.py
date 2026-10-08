import io
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from app.parser import parse_workbook

RED = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
GREEN = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")
GRAY = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
EXAMPLE = Path(__file__).resolve().parents[2] / "example.xlsx"


def build_sample() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "2414-2415"
    sheet["A1"] = "Расписание занятий на 1 семестр 2026-2027 учебный год"
    sheet["C2"] = "ТУР-2415/1"
    sheet.merge_cells("C2:D2")
    sheet["E2"] = "ДО-2414/2"
    sheet["G2"] = "ТМ-261"

    sheet["A3"] = "Понедельник"
    sheet.merge_cells("A3:A6")
    sheet["B3"] = 1
    sheet.merge_cells("B3:B4")
    sheet.merge_cells("C3:C4")
    sheet["C3"] = "Иностранный язык\nКидалова О.В Цветкова А.В."
    sheet["D3"] = "502\n204"
    sheet["E3"] = "Физика\nИванов И.И."
    sheet["E3"].fill = RED
    sheet["F3"] = 313
    sheet["F3"].fill = RED
    sheet["E4"] = "Химия\nПетров П.П."
    sheet["E4"].fill = GREEN
    sheet["F4"] = "с/з"
    sheet["F4"].fill = GREEN

    sheet["B5"] = 2
    sheet.merge_cells("B5:B6")
    sheet["C5"] = "История\nБурич П.А."
    sheet["C5"].fill = RED
    sheet["D5"] = 119
    sheet["D5"].fill = RED
    sheet["E5"] = "Учебная практика"
    sheet["E5"].fill = GRAY
    sheet["E6"] = "Башилова Ю.А."
    sheet["E6"].fill = GRAY
    sheet["F5"] = 316
    sheet["G6"] = "ОБЗР\nКулибов М.Ю."
    sheet["G6"].fill = GREEN
    sheet["H6"] = 10
    sheet["H6"].fill = GREEN

    sheet["B7"] = 3
    sheet.merge_cells("B7:B8")
    sheet["C7"] = "Информатика"
    sheet["C8"] = "Проценко Н.В Леонов А.С"
    sheet["D7"] = 318
    sheet["D8"] = 403
    sheet["G7"] = "Разработка мобильных приложений\nФедоров В.В."
    sheet["G8"] = "Федоров В.В."

    empty = book.create_sheet("пусто")
    empty["A1"] = "нет групп"

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


class ParserTest(unittest.TestCase):
    def test_groups_rooms_and_weeks(self):
        result = parse_workbook(io.BytesIO(build_sample()))

        self.assertIn("Расписание занятий", result.title)
        self.assertEqual(result.warnings, ["лист «пусто»: не найдены группы"])

        def pick(group, week, pair):
            return [
                lesson
                for lesson in result.lessons
                if lesson.group_name == group and lesson.week == week and lesson.pair == pair
            ]

        names = {lesson.group_name for lesson in result.lessons}
        self.assertEqual(names, {"ТУР-2415/1", "ДО-2414/2", "ТМ-261"})

        language_green = pick("ТУР-2415/1", "green", 1)
        language_red = pick("ТУР-2415/1", "red", 1)
        self.assertEqual(len(language_green), 1)
        self.assertEqual(len(language_red), 1)
        self.assertEqual(language_green[0].subject, "Иностранный язык")
        self.assertEqual(language_green[0].teacher, "Кидалова О.В Цветкова А.В.")
        self.assertEqual(language_green[0].room, "502, 204")
        self.assertEqual(language_green[0].day, "Понедельник")
        self.assertEqual(language_red[0].room, "502, 204")

        physics = pick("ДО-2414/2", "red", 1)
        self.assertEqual(len(physics), 1)
        self.assertEqual(physics[0].subject, "Физика")
        self.assertEqual(physics[0].teacher, "Иванов И.И.")
        self.assertEqual(physics[0].room, "313")
        self.assertEqual(pick("ДО-2414/2", "green", 1)[0].subject, "Химия")
        self.assertEqual(pick("ДО-2414/2", "green", 1)[0].room, "с/з")

        history = pick("ТУР-2415/1", "red", 2)
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0].subject, "История")
        self.assertEqual(history[0].teacher, "Бурич П.А.")
        self.assertEqual(history[0].room, "119")
        self.assertEqual(pick("ТУР-2415/1", "green", 2), [])

        obzr = pick("ТМ-261", "green", 2)
        self.assertEqual(len(obzr), 1)
        self.assertEqual(obzr[0].subject, "ОБЗР")
        self.assertEqual(obzr[0].teacher, "Кулибов М.Ю.")
        self.assertEqual(obzr[0].room, "10")
        self.assertEqual(pick("ТМ-261", "red", 2), [])

        practice_green = pick("ДО-2414/2", "green", 2)
        practice_red = pick("ДО-2414/2", "red", 2)
        self.assertEqual(len(practice_green), 1)
        self.assertEqual(len(practice_red), 1)
        self.assertEqual(practice_green[0].subject, "Учебная практика")
        self.assertEqual(practice_green[0].teacher, "Башилова Ю.А.")
        self.assertEqual(practice_green[0].room, "316")
        self.assertEqual(practice_red[0].teacher, "Башилова Ю.А.")

        informatics_green = pick("ТУР-2415/1", "green", 3)
        informatics_red = pick("ТУР-2415/1", "red", 3)
        self.assertEqual(len(informatics_green), 1)
        self.assertEqual(len(informatics_red), 1)
        self.assertEqual(informatics_green[0].subject, "Информатика")
        self.assertEqual(informatics_green[0].teacher, "Проценко Н.В Леонов А.С")
        self.assertEqual(informatics_green[0].room, "318, 403")

        mobile = pick("ТМ-261", "green", 3)
        self.assertEqual(len(mobile), 1)
        self.assertEqual(len(pick("ТМ-261", "red", 3)), 1)
        self.assertEqual(mobile[0].subject, "Разработка мобильных приложений")
        self.assertEqual(mobile[0].teacher, "Федоров В.В.")

    @unittest.skipUnless(EXAMPLE.exists(), "нет example.xlsx")
    def test_example_file(self):
        result = parse_workbook(EXAMPLE)
        self.assertEqual(result.warnings, [])

        def one(group, week, pair, day="Понедельник"):
            found = [
                lesson
                for lesson in result.lessons
                if lesson.group_name == group
                and lesson.week == week
                and lesson.pair == pair
                and lesson.day == day
            ]
            self.assertEqual(len(found), 1, (group, week, pair, day, found))
            return found[0]

        databases = one("ИС-2411/2", "green", 1)
        self.assertEqual(databases.subject, "Технология разработки и защиты баз данных")
        self.assertEqual(databases.teacher, "Иванова Н.Л")
        self.assertEqual(databases.room, "319")
        red_databases = one("ИС-2411/2", "red", 1)
        self.assertEqual(red_databases.subject, databases.subject)
        self.assertEqual(red_databases.teacher, databases.teacher)

        practice = one("ЮР-2413/1", "red", 1)
        self.assertEqual(practice.subject, "Учебная практика")
        self.assertEqual(practice.teacher, "Башилова Ю.А")
        self.assertEqual(practice.room, "316")
        self.assertEqual(one("ЮР-2413/1", "green", 1).teacher, "Башилова Ю.А")

        language = one("ЮР-2413/2", "green", 1)
        self.assertEqual(language.subject, "Иностранный язык")
        self.assertEqual(language.teacher, "Кидалова О.В Цветкова А.В.")
        self.assertEqual(language.room, "502, 204")
        self.assertEqual(
            [
                lesson
                for lesson in result.lessons
                if lesson.group_name == "ЮР-2413/2"
                and lesson.week == "red"
                and lesson.pair == 1
                and lesson.day == "Понедельник"
            ],
            [],
        )

        philosophy = one("ИС-2411/2", "red", 2)
        self.assertEqual(philosophy.subject, "Основы философии")
        self.assertEqual(philosophy.teacher, "Божко Ю.В.")
        self.assertEqual(philosophy.room, "117")
        psychology = one("ИС-2411/2", "green", 2)
        self.assertEqual(psychology.subject, "Психология общения")
        self.assertEqual(psychology.teacher, "Стрельцова В.В.")
        self.assertEqual(psychology.room, "118")

        for week in ("red", "green"):
            sport = one("ЮР-2413/2", week, 4)
            self.assertEqual(sport.subject, "Физическая культура")
            self.assertEqual(sport.room, "с/з")

        labor = one("ЮР-2413/1", "green", 4)
        self.assertEqual(labor.subject, "Трудовое право")
        self.assertEqual(labor.teacher, "Башилова Ю.А.")
        self.assertEqual(labor.room, "316")
        self.assertEqual(one("ЮР-2413/1", "red", 4).subject, "Трудовое право")

        mobile = one("ИС-2411/2", "red", 5)
        self.assertEqual(mobile.subject, "Разработка мобильных приложений")
        self.assertEqual(mobile.teacher, "Федоров В.В.")
        self.assertEqual(one("ИС-2411/2", "green", 5).teacher, "Федоров В.В.")


if __name__ == "__main__":
    unittest.main()
