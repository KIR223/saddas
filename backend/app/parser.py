"""Чтение расписания из xlsx: группы в шапке, неделя по цвету внутри пары."""

import colorsys
import re
from dataclasses import dataclass
from xml.etree import ElementTree as ET

from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.worksheet.worksheet import Worksheet

GROUP_RE = re.compile(r"^[А-ЯЁA-Z]{1,6}-\d{2,4}(?:/\d{1,2})?$")
ROOM_RE = re.compile(r"^(?:\d{1,4}[A-ZА-ЯЁ]?|[СC]/[ЗZ])$", re.IGNORECASE)
THEME_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

DAY_ORDER = [
    "Понедельник",
    "Вторник",
    "Среда",
    "Четверг",
    "Пятница",
    "Суббота",
    "Воскресенье",
]


@dataclass
class Lesson:
    group_name: str
    day: str
    pair: int
    week: str
    subject: str
    teacher: str
    room: str


@dataclass
class ParseResult:
    title: str
    lessons: list[Lesson]
    warnings: list[str]


@dataclass
class PairSlot:
    day: str
    pair: int
    rows: list[int]


def parse_workbook(source) -> ParseResult:
    from openpyxl import load_workbook

    workbook = load_workbook(source, data_only=False)
    theme = load_theme_colors(workbook)
    lessons: list[Lesson] = []
    warnings: list[str] = []
    title = ""
    for sheet in workbook.worksheets:
        sheet_title = find_title(sheet)
        if sheet_title and not title:
            title = sheet_title
        parsed, sheet_warnings = parse_sheet(sheet, theme)
        lessons.extend(parsed)
        warnings.extend(sheet_warnings)
    return ParseResult(title=title, lessons=lessons, warnings=warnings)


def parse_sheet(sheet: Worksheet, theme: list[str]) -> tuple[list[Lesson], list[str]]:
    origins = merge_origins(sheet)
    header = find_group_row(sheet, origins)
    if header is None:
        return [], [f"лист «{sheet.title}»: не найдены группы"]
    header_row, groups = header
    groups = subject_columns(groups, origins, header_row)
    if not groups:
        return [], [f"лист «{sheet.title}»: не найдены группы"]
    first_group_col = min(groups)
    lessons: list[Lesson] = []
    for slot in pair_slots(sheet, origins, header_row, first_group_col, groups):
        for col, group_name in groups.items():
            lessons.extend(lessons_for_group(sheet, origins, slot, col, group_name, groups, theme))
    return lessons, []


def subject_columns(groups: dict[int, str], origins: dict, header_row: int) -> dict[int, str]:
    """Правый столбец объединённого заголовка — аудитория, не вторая группа."""
    skipped: set[int] = set()
    for col in sorted(groups):
        right = col + 1
        if right not in groups or groups[col] != groups[right]:
            continue
        left_origin = origins.get((header_row, col), (header_row, col))
        right_origin = origins.get((header_row, right), (header_row, right))
        if left_origin == right_origin:
            skipped.add(right)
    return {col: name for col, name in groups.items() if col not in skipped}


def pair_slots(
    sheet: Worksheet,
    origins: dict,
    header_row: int,
    first_group_col: int,
    groups: dict[int, str],
) -> list[PairSlot]:
    slots: list[PairSlot] = []
    current_day = None
    current: PairSlot | None = None
    for row in range(header_row + 1, (sheet.max_row or header_row) + 1):
        day, pair = day_and_pair(sheet, row, first_group_col, origins)
        if day:
            current_day = day
        if not current_day:
            continue
        adjacent = current is not None and row == current.rows[-1] + 1 and current.day == current_day
        if pair is not None and adjacent and current.pair == pair:
            current.rows.append(row)
            continue
        if pair is not None:
            current = PairSlot(day=current_day, pair=pair, rows=[row])
            slots.append(current)
            continue
        if adjacent and row_has_lesson(sheet, row, groups, origins):
            current.rows.append(row)
    return slots


def lessons_for_group(
    sheet: Worksheet,
    origins: dict,
    slot: PairSlot,
    col: int,
    group_name: str,
    groups: dict[int, str],
    theme: list[str],
) -> list[Lesson]:
    room_col = col + 1 if col + 1 not in groups else None
    rows = slot.rows
    if not rows:
        return []
    # Одна объединённая ячейка на обе строки пары: цвет решает, одна это неделя или обе.
    if len(rows) > 1 and cell_covers_rows(origins, rows, col):
        origin = origins.get((rows[0], col), (rows[0], col))
        return placed_lesson(
            slot,
            group_name,
            split_subject_teacher(cell_text(sheet, rows[0], col, origins)),
            read_room(sheet, rows, room_col, origins),
            weeks_of(fill_kind(sheet.cell(origin[0], origin[1]), theme)),
        )
    owned: list[tuple[int, str]] = []
    for index, row in enumerate(rows[:2]):
        if origins.get((row, col), (row, col)) != (row, col):
            continue
        if not cell_text(sheet, row, col, origins):
            continue
        owned.append((row, "red" if index == 0 else "green"))
    if not owned:
        return []
    # Без красной и зелёной заливки это одно занятие на обе недели:
    # предмет в верхней строке, преподаватель в нижней.
    if all(fill_kind(sheet.cell(row, col), theme) == "neutral" for row, _ in owned):
        subject, teacher = combine_both_weeks([cell_text(sheet, row, col, origins) for row, _ in owned])
        return placed_lesson(
            slot,
            group_name,
            (subject, teacher),
            read_room(sheet, [row for row, _ in owned], room_col, origins),
            ("green", "red"),
        )
    found: list[Lesson] = []
    for row, default_week in owned:
        kind = fill_kind(sheet.cell(row, col), theme)
        week = kind if kind in ("red", "green") else default_week
        found.extend(
            placed_lesson(
                slot,
                group_name,
                split_subject_teacher(cell_text(sheet, row, col, origins)),
                read_room(sheet, [row], room_col, origins),
                (week,),
            )
        )
    return found


def placed_lesson(
    slot: PairSlot,
    group_name: str,
    subject_teacher: tuple[str, str],
    room: str,
    weeks: tuple[str, ...],
) -> list[Lesson]:
    subject, teacher = subject_teacher
    if not subject:
        return []
    return [
        Lesson(
            group_name=group_name,
            day=slot.day,
            pair=slot.pair,
            week=week,
            subject=subject,
            teacher=teacher,
            room=room,
        )
        for week in weeks
    ]


def weeks_of(kind: str) -> tuple[str, ...]:
    if kind in ("red", "green"):
        return (kind,)
    return ("green", "red")


def combine_both_weeks(parts: list[str]) -> tuple[str, str]:
    blocks = [text_lines(part) for part in parts]
    blocks = [block for block in blocks if block]
    if not blocks:
        return "", ""
    subject = blocks[0][0]
    teacher_parts: list[str] = []
    for block in blocks:
        for line in block:
            if line == subject or line in teacher_parts:
                continue
            teacher_parts.append(line)
    return subject, " ".join(teacher_parts)


def text_lines(text: str) -> list[str]:
    return [part.strip() for part in re.split(r"\r?\n", text) if part.strip()]


def cell_covers_rows(origins: dict, rows: list[int], col: int) -> bool:
    if len(rows) <= 1:
        return True
    origin = origins.get((rows[0], col), (rows[0], col))
    return all(origins.get((row, col), (row, col)) == origin for row in rows)


def row_has_lesson(sheet: Worksheet, row: int, groups: dict[int, str], origins: dict) -> bool:
    return any(cell_text(sheet, row, col, origins) for col in groups)


def find_group_row(sheet: Worksheet, origins: dict) -> tuple[int, dict[int, str]] | None:
    limit = min(sheet.max_row or 1, 30)
    candidates: list[tuple[int, dict[int, str]]] = []
    for row in range(1, limit + 1):
        found: dict[int, str] = {}
        for col in range(1, (sheet.max_column or 1) + 1):
            name = normalize_group(cell_value(sheet, row, col, origins))
            if name:
                found[col] = name
        if found:
            candidates.append((row, found))
    if not candidates:
        return None
    best = max(len(found) for _, found in candidates)
    for row, found in candidates:
        if len(found) == best:
            return row, found
    return None


def find_title(sheet: Worksheet) -> str:
    max_row = min(sheet.max_row or 1, 8)
    max_col = min(sheet.max_column or 1, 8)
    for row in sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cell in row:
            if cell.value and "расписан" in str(cell.value).lower():
                return " ".join(str(cell.value).split())
    return ""


def day_and_pair(sheet: Worksheet, row: int, first_group_col: int, origins: dict) -> tuple[str | None, int | None]:
    day = None
    pair = None
    for col in range(1, first_group_col):
        value = cell_value(sheet, row, col, origins)
        day_name = normalize_day(value)
        if day_name:
            day = day_name
            continue
        pair_number = parse_pair(value)
        if pair_number is not None:
            pair = pair_number
    return day, pair


def normalize_group(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    text = text.replace("–", "-").replace("—", "-").replace("−", "-")
    text = re.sub(r"\s+", "", text).upper()
    if GROUP_RE.fullmatch(text):
        return text
    return None


def normalize_day(value) -> str | None:
    if value is None:
        return None
    compact = re.sub(r"\s+", "", str(value).casefold())
    if not compact:
        return None
    for name in DAY_ORDER:
        if name.casefold() in compact:
            return name
    return None


def parse_pair(value) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        if float(value) != int(value):
            return None
        number = int(value)
    else:
        text = str(value).strip()
        if not re.fullmatch(r"\d{1,2}", text):
            return None
        number = int(text)
    if 1 <= number <= 12:
        return number
    return None


def fill_kind(cell, theme: list[str]) -> str:
    fill = cell.fill
    if fill is None or fill.patternType in (None, "none"):
        return "neutral"
    rgb = color_to_rgb(fill.fgColor, theme)
    if rgb is None:
        return "neutral"
    return classify_rgb(*rgb)


def classify_rgb(red: int, green: int, blue: int) -> str:
    # Серый и белый — обе недели. Порог 18, а не выше: бледно-зелёный Excel
    # (226, 240, 217) отличается от серого всего на 23 пункта.
    if max(red, green, blue) - min(red, green, blue) < 18:
        return "neutral"
    if green >= red + 10 and green >= blue:
        return "green"
    if red >= green + 15 and red >= blue:
        return "red"
    return "neutral"


def color_to_rgb(color, theme: list[str]) -> tuple[int, int, int] | None:
    if color is None or not color.type:
        return None
    hex_rgb = None
    if color.type == "rgb" and color.rgb:
        hex_rgb = str(color.rgb)
        if hex_rgb in ("00000000", "0"):
            return None
    elif color.type == "indexed" and color.indexed is not None:
        hex_rgb = indexed_hex(color.indexed)
    elif color.type == "theme" and color.theme is not None:
        if color.theme < 0 or color.theme >= len(theme):
            return None
        hex_rgb = apply_tint(theme[color.theme], float(color.tint or 0))
    if not hex_rgb:
        return None
    hex_rgb = hex_rgb[-6:]
    try:
        return tuple(int(hex_rgb[i : i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return None


def indexed_hex(index: int) -> str | None:
    try:
        value = COLOR_INDEX[index]
    except (IndexError, TypeError):
        return None
    if not value:
        return None
    return str(value)


def apply_tint(hex_rgb: str, tint: float) -> str:
    base = hex_rgb[-6:]
    if not tint:
        return base
    red, green, blue = [int(base[i : i + 2], 16) / 255 for i in (0, 2, 4)]
    hue, light, sat = colorsys.rgb_to_hls(red, green, blue)
    if tint < 0:
        light = light * (1 + tint)
    else:
        light = light * (1 - tint) + tint
    red, green, blue = colorsys.hls_to_rgb(hue, max(0, min(1, light)), sat)
    return f"{round(red * 255):02X}{round(green * 255):02X}{round(blue * 255):02X}"


def load_theme_colors(workbook) -> list[str]:
    raw = getattr(workbook, "loaded_theme", None)
    if not raw:
        return []
    try:
        root = ET.fromstring(raw)
    except ET.ParseError:
        return []
    scheme = root.find(".//a:clrScheme", THEME_NS)
    if scheme is None:
        return []
    colors: list[str] = []
    for child in list(scheme):
        srgb = child.find("a:srgbClr", THEME_NS)
        if srgb is not None and srgb.get("val"):
            colors.append(srgb.get("val"))
            continue
        sys_color = child.find("a:sysClr", THEME_NS)
        if sys_color is not None and sys_color.get("lastClr"):
            colors.append(sys_color.get("lastClr"))
            continue
        colors.append("000000")
    return colors


def split_subject_teacher(text: str) -> tuple[str, str]:
    lines = [part.strip() for part in re.split(r"\r?\n", text) if part.strip()]
    if not lines:
        return "", ""
    return lines[0], " ".join(lines[1:])


def read_room(sheet: Worksheet, rows: list[int], col: int | None, origins: dict) -> str:
    if col is None:
        return ""
    found: list[str] = []
    for row in rows:
        for line in re.split(r"\r?\n", cell_text(sheet, row, col, origins)):
            room = line.strip()
            if room and ROOM_RE.fullmatch(room) and room not in found:
                found.append(room)
    return ", ".join(found)


def cell_text(sheet: Worksheet, row: int, col: int, origins: dict) -> str:
    value = cell_value(sheet, row, col, origins)
    if value is None:
        return ""
    return str(value).strip()


def merge_origins(sheet: Worksheet) -> dict[tuple[int, int], tuple[int, int]]:
    origins: dict[tuple[int, int], tuple[int, int]] = {}
    for merged in sheet.merged_cells.ranges:
        origin = (merged.min_row, merged.min_col)
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                origins[(row, col)] = origin
    return origins


def cell_value(sheet: Worksheet, row: int, col: int, origins: dict):
    src_row, src_col = origins.get((row, col), (row, col))
    return sheet.cell(src_row, src_col).value
